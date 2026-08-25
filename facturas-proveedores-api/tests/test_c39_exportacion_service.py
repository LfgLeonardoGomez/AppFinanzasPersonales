"""
C-39, task group 6 — `ExportacionCuentaCorrienteService` (design.md D1).

Service-layer tests against real Postgres (testcontainers). Mirrors the
patterns of `test_cuenta_corriente_service.py` and
`test_c35_cuenta_corriente_cliente_service.py`.

Covers:
- Structural guard (task 6.1): the module never imports a repository nor
  the FIFO engine — it only calls the cuenta-corriente services that
  already exist.
- The test that matters (task 6.2): exporting an account and reading its
  cuenta-corriente endpoint give the SAME saldo and the SAME movements.
- 404 on foreign / missing supplier and customer (task 6.3, 6.4) — never
  403, never distinguishable from each other.
- Ownership verification runs BEFORE any byte is generated (task 6.5).
"""

import ast
import io
import uuid
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlmodel import Session, SQLModel

import app.models  # noqa: F401

from tests.conftest import crear_negocio


@pytest.fixture(scope="module")
def engine(db_url: str):
    eng = create_engine(db_url, echo=False)
    SQLModel.metadata.create_all(eng)
    yield eng
    eng.dispose()


@pytest.fixture
def session(engine):
    with Session(engine) as s:
        yield s
        s.rollback()


# ── Factories ────────────────────────────────────────────────────────────────


def _make_usuario(session: Session, negocio_id=None):
    from app.models.usuario import Usuario
    from app.core.uuid_utils import new_uuid

    u = Usuario(
        negocio_id=negocio_id or crear_negocio(session).id,
        id=new_uuid(),
        email=f"exp_{uuid.uuid4().hex[:8]}@test.com",
        nombre="Export Test",
        password_hash="$argon2id$v=19$m=65536,t=3,p=4$fakehash",
    )
    session.add(u)
    session.flush()
    return u


def _make_proveedor(session: Session, negocio_id: uuid.UUID, deleted: bool = False):
    from app.models.proveedor import Proveedor
    from app.core.uuid_utils import new_uuid
    from app.models.enums import CategoriaProveedor

    p = Proveedor(
        id=new_uuid(),
        negocio_id=negocio_id,
        nombre=f"Prov {uuid.uuid4().hex[:6]}",
        categoria=CategoriaProveedor.OTRO,
    )
    if deleted:
        p.deleted_at = datetime.now(timezone.utc)
    session.add(p)
    session.flush()
    return p


def _make_factura_db(session, negocio_id, proveedor_id, monto_total, fecha_emision):
    from app.models.factura import Factura
    from app.core.uuid_utils import new_uuid
    from app.models.enums import OrigenDocumento

    f = Factura(
        id=new_uuid(),
        negocio_id=negocio_id,
        proveedor_id=proveedor_id,
        fecha_emision=fecha_emision,
        monto_total=monto_total,
        origen=OrigenDocumento.MANUAL,
    )
    session.add(f)
    session.flush()
    return f


def _make_pago_db(session, negocio_id, proveedor_id, monto, fecha):
    from app.models.pago import Pago
    from app.core.uuid_utils import new_uuid
    from app.models.enums import MetodoPago, OrigenDocumento

    p = Pago(
        id=new_uuid(),
        negocio_id=negocio_id,
        proveedor_id=proveedor_id,
        monto=monto,
        fecha=fecha,
        metodo=MetodoPago.EFECTIVO,
        origen=OrigenDocumento.MANUAL,
    )
    session.add(p)
    session.flush()
    return p


def _make_cliente(session: Session, negocio_id: uuid.UUID, deleted: bool = False):
    from app.models.cliente import Cliente

    c = Cliente(
        negocio_id=negocio_id,
        nombre=f"Cliente {uuid.uuid4().hex[:6]}",
        nombre_normalizado=f"cliente {uuid.uuid4().hex[:8]}",
    )
    if deleted:
        c.deleted_at = datetime.now(timezone.utc)
    session.add(c)
    session.flush()
    return c


# ── 6.1 — Structural guard ───────────────────────────────────────────────────


class TestNoAbreQueriesPropias:
    def test_no_importa_repositorios_ni_el_motor_fifo(self):
        """
        Espíritu del guard de C-28 (task 6.1): el módulo del service de
        exportación no debe importar `app.repositories.*` ni
        `cuenta_corriente_engine` — solo puede apoyarse en
        ProveedorService/ClienteService, que ya hacen esa autorización y ese
        cálculo.
        """
        path = (
            Path(__file__).resolve().parents[1]
            / "app"
            / "services"
            / "exportacion_cuenta_corriente_service.py"
        )
        tree = ast.parse(path.read_text(encoding="utf-8"))
        modulos_importados: set[str] = set()
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                for alias in node.names:
                    modulos_importados.add(alias.name)
            elif isinstance(node, ast.ImportFrom) and node.module:
                modulos_importados.add(node.module)

        for modulo in modulos_importados:
            assert not modulo.startswith("app.repositories"), (
                f"exportacion_cuenta_corriente_service.py importa {modulo} — "
                f"D1 prohíbe abrir queries propias."
            )
            assert "cuenta_corriente_engine" not in modulo, (
                f"exportacion_cuenta_corriente_service.py importa {modulo} — "
                f"D1 prohíbe importar el motor FIFO."
            )


# ── 6.2 — El test que importa ────────────────────────────────────────────────


class TestCoincideConLaPantalla:
    def test_proveedor_export_coincide_con_get_cuenta_corriente(self, session: Session):
        from app.services.proveedor_service import ProveedorService
        from app.services.exportacion_cuenta_corriente_service import (
            ExportacionCuentaCorrienteService,
        )

        u = _make_usuario(session)
        p = _make_proveedor(session, u.negocio_id)
        _make_factura_db(session, u.negocio_id, p.id, Decimal("1000.00"), date.today())
        _make_pago_db(session, u.negocio_id, p.id, Decimal("300.00"), date.today())
        session.commit()

        esperado = ProveedorService(session).get_cuenta_corriente(u.negocio_id, p.id)

        archivo = ExportacionCuentaCorrienteService(session).exportar_proveedor(
            u.negocio_id, p.id, formato="xlsx", incluir_historial=True,
        )
        wb = load_workbook(io.BytesIO(archivo.contenido))
        valores = [
            c.value for row in wb.active.iter_rows() for c in row if c.value is not None
        ]

        assert float(esperado.saldo) in valores
        assert len(esperado.historial) == 2
        for h in esperado.historial:
            assert float(h["monto"]) in valores
            assert float(h["saldo_acumulado"]) in valores

    def test_cliente_export_coincide_con_get_cuenta_corriente(self, session: Session):
        from app.services.cliente_service import ClienteService
        from app.models.venta import Venta
        from app.models.enums import FormaPago
        from app.services.exportacion_cuenta_corriente_service import (
            ExportacionCuentaCorrienteService,
        )

        u = _make_usuario(session)
        c = _make_cliente(session, u.negocio_id)
        venta = Venta(
            negocio_id=u.negocio_id, cliente_id=c.id, monto=Decimal("500.00"),
            fecha=date.today(), forma_pago=FormaPago.CUENTA_CORRIENTE,
        )
        session.add(venta)
        session.flush()
        session.commit()

        esperado = ClienteService(session).get_cuenta_corriente(u.negocio_id, c.id)

        archivo = ExportacionCuentaCorrienteService(session).exportar_cliente(
            u.negocio_id, c.id, formato="xlsx", incluir_historial=True,
        )
        wb = load_workbook(io.BytesIO(archivo.contenido))
        valores = [
            v.value for row in wb.active.iter_rows() for v in row if v.value is not None
        ]

        assert float(esperado.saldo) in valores
        assert len(esperado.historial) == 1
        assert float(esperado.historial[0]["monto"]) in valores


# ── 6.3 / 6.4 — 404, nunca 403, indistinguible ───────────────────────────────


class TestAislamientoPorNegocio:
    def test_proveedor_de_otro_negocio_da_404(self, session: Session):
        from app.services.exportacion_cuenta_corriente_service import (
            ExportacionCuentaCorrienteService,
        )

        u_a = _make_usuario(session)
        u_b = _make_usuario(session)
        prov_b = _make_proveedor(session, u_b.negocio_id)
        session.commit()

        svc = ExportacionCuentaCorrienteService(session)
        with pytest.raises(HTTPException) as exc_info:
            svc.exportar_proveedor(u_a.negocio_id, prov_b.id, formato="pdf")
        assert exc_info.value.status_code == 404

    def test_cliente_de_otro_negocio_da_404(self, session: Session):
        from app.services.exportacion_cuenta_corriente_service import (
            ExportacionCuentaCorrienteService,
        )

        u_a = _make_usuario(session)
        u_b = _make_usuario(session)
        cli_b = _make_cliente(session, u_b.negocio_id)
        session.commit()

        svc = ExportacionCuentaCorrienteService(session)
        with pytest.raises(HTTPException) as exc_info:
            svc.exportar_cliente(u_a.negocio_id, cli_b.id, formato="pdf")
        assert exc_info.value.status_code == 404

    def test_proveedor_inexistente_da_404_indistinguible(self, session: Session):
        from app.services.exportacion_cuenta_corriente_service import (
            ExportacionCuentaCorrienteService,
        )

        u = _make_usuario(session)
        session.commit()

        svc = ExportacionCuentaCorrienteService(session)
        with pytest.raises(HTTPException) as exc_info:
            svc.exportar_proveedor(u.negocio_id, uuid.uuid4(), formato="pdf")
        assert exc_info.value.status_code == 404


# ── 6.5 — la verificación corre antes de generar nada ────────────────────────


class TestVerificacionAntesDeGenerar:
    def test_no_llama_a_los_generadores_si_el_recurso_es_ajeno(self, session: Session, monkeypatch):
        import app.services.exportacion_cuenta_corriente_service as mod

        def _explota(*args, **kwargs):
            raise AssertionError("no debería generarse ningún documento para un recurso ajeno")

        monkeypatch.setattr(mod, "generar_pdf", _explota)
        monkeypatch.setattr(mod, "generar_xlsx", _explota)

        u_a = _make_usuario(session)
        u_b = _make_usuario(session)
        prov_b = _make_proveedor(session, u_b.negocio_id)
        session.commit()

        svc = mod.ExportacionCuentaCorrienteService(session)
        with pytest.raises(HTTPException) as exc_info:
            svc.exportar_proveedor(u_a.negocio_id, prov_b.id, formato="pdf")
        assert exc_info.value.status_code == 404


# ── Contradicción incluir_historial=False + desde/hasta ──────────────────────


class TestParametrosContradictorios:
    def test_rango_sin_incluir_historial_da_422(self, session: Session):
        from app.services.exportacion_cuenta_corriente_service import (
            ExportacionCuentaCorrienteService,
        )

        u = _make_usuario(session)
        p = _make_proveedor(session, u.negocio_id)
        session.commit()

        svc = ExportacionCuentaCorrienteService(session)
        with pytest.raises(HTTPException) as exc_info:
            svc.exportar_proveedor(
                u.negocio_id, p.id, formato="pdf",
                incluir_historial=False, desde=date(2026, 1, 1),
            )
        assert exc_info.value.status_code == 422


# ── Tope (integración con el service) ────────────────────────────────────────


class TestTopeIntegrado:
    def test_excede_el_tope_da_422_y_no_genera_nada(self, session: Session, monkeypatch):
        import app.services.exportacion_cuenta_corriente_service as mod

        def _explota(*args, **kwargs):
            raise AssertionError("no debería generarse ningún documento por sobre el tope")

        monkeypatch.setattr(mod, "generar_pdf", _explota)
        monkeypatch.setitem(mod.TOPE_FILAS, "pdf", 1)

        u = _make_usuario(session)
        p = _make_proveedor(session, u.negocio_id)
        _make_factura_db(session, u.negocio_id, p.id, Decimal("100.00"), date(2026, 1, 1))
        _make_pago_db(session, u.negocio_id, p.id, Decimal("50.00"), date(2026, 1, 2))
        session.commit()

        svc = mod.ExportacionCuentaCorrienteService(session)
        with pytest.raises(HTTPException) as exc_info:
            svc.exportar_proveedor(u.negocio_id, p.id, formato="pdf", incluir_historial=True)
        assert exc_info.value.status_code == 422
        assert exc_info.value.detail["cantidad_movimientos"] == 2
