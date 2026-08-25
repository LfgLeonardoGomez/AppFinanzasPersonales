"""
C-39, task group 4 — PDF/XLSX generation and filename (design.md D3, D6).

Consumes `DocumentoExportacion` from `exportacion_armador` and produces
bytes. Neither generator recomputes anything — they only render what the
armador already built (D1).
"""

import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app.services.exportacion_armador import (
    DocumentoExportacion,
    FilaDocumento,
    SaldoRotulado,
    armar_documento,
)
from app.services.exportacion_generadores import (
    construir_nombre_archivo,
    generar_pdf,
    generar_xlsx,
)


def _fila(fecha, tipo, monto, saldo_acumulado, id_=None):
    import uuid

    return {
        "id": id_ or uuid.uuid4(),
        "tipo": tipo,
        "fecha": date.fromisoformat(fecha),
        "monto": Decimal(monto),
        "saldo_acumulado": Decimal(saldo_acumulado),
        "archivo_url": None,
    }


HISTORIAL = [
    _fila("2026-01-05", "FACTURA", "1000.00", "1000.00"),
    _fila("2026-01-10", "PAGO", "300.00", "700.00"),
]


def _documento(**overrides) -> DocumentoExportacion:
    kwargs = dict(
        saldo=Decimal("700.00"),
        historial=HISTORIAL,
        nombre_cuenta="Ferretería El Tornillo",
        nombre_negocio="Kiosco Don José",
        fecha_emision=date(2026, 8, 25),
        incluir_historial=False,
    )
    kwargs.update(overrides)
    return armar_documento(**kwargs)


# ── XLSX ─────────────────────────────────────────────────────────────────────


class TestGenerarXlsx:
    def test_produce_archivo_reabrible_con_los_montos_del_armador(self):
        doc = _documento(incluir_historial=True)
        contenido = generar_xlsx(doc)

        wb = load_workbook(io.BytesIO(contenido))
        ws = wb.active
        valores = [cell.value for row in ws.iter_rows() for cell in row if cell.value is not None]

        assert "Ferretería El Tornillo" in valores
        # Montos del armador presentes como número
        assert 1000.0 in valores
        assert 300.0 in valores
        assert 700.0 in valores  # saldo del encabezado

    def test_triangulacion_sin_historial_vs_con_historial_distinta_cantidad_de_filas(self):
        doc_sin = _documento(incluir_historial=False)
        doc_con = _documento(incluir_historial=True)

        contenido_sin = generar_xlsx(doc_sin)
        contenido_con = generar_xlsx(doc_con)

        ws_sin = load_workbook(io.BytesIO(contenido_sin)).active
        ws_con = load_workbook(io.BytesIO(contenido_con)).active

        assert ws_con.max_row > ws_sin.max_row
        # El documento sin historial igual trae el encabezado (nombre del negocio).
        valores_sin = [cell.value for row in ws_sin.iter_rows() for cell in row if cell.value is not None]
        assert "Kiosco Don José" in valores_sin

    def test_con_rango_incluye_la_fila_de_saldo_anterior(self):
        doc = _documento(
            incluir_historial=True,
            desde=date(2026, 1, 10),
            hasta=date(2026, 1, 10),
        )
        contenido = generar_xlsx(doc)
        ws = load_workbook(io.BytesIO(contenido)).active
        valores = [cell.value for row in ws.iter_rows() for cell in row if cell.value is not None]
        assert "Saldo anterior" in valores
        assert 1000.0 in valores  # saldo_anterior calculado por el armador


# ── PDF ──────────────────────────────────────────────────────────────────────


class TestGenerarPdf:
    def test_produce_pdf_valido_con_nombre_de_cuenta_y_saldo_en_el_texto(self):
        doc = _documento(incluir_historial=False)
        contenido = generar_pdf(doc)

        assert contenido.startswith(b"%PDF-")
        assert len(contenido) > 0

        reader = PdfReader(io.BytesIO(contenido))
        texto = "\n".join(page.extract_text() for page in reader.pages)
        assert "Ferretería El Tornillo" in texto or "Ferreteria El Tornillo" in texto
        assert "700" in texto  # saldo

    def test_con_rango_el_pdf_tambien_muestra_saldo_anterior(self):
        doc = _documento(
            incluir_historial=True,
            desde=date(2026, 1, 10),
            hasta=date(2026, 1, 10),
        )
        contenido = generar_pdf(doc)
        reader = PdfReader(io.BytesIO(contenido))
        texto = "\n".join(page.extract_text() for page in reader.pages)
        assert "anterior" in texto.lower()


# ── Nombre de archivo (D6, task 4.6) ──────────────────────────────────────────


class TestConstruirNombreArchivo:
    def test_incluye_tipo_nombre_y_fecha(self):
        nombre = construir_nombre_archivo("Ferretería El Tornillo", date(2026, 8, 25), "pdf")
        assert nombre.startswith("cuenta-corriente-")
        assert "2026-08-25" in nombre
        assert nombre.endswith(".pdf")

    def test_triangulacion_acentos_y_espacios_se_normalizan(self):
        nombre = construir_nombre_archivo("Ñandú & Cía. S.A.", date(2026, 1, 1), "xlsx")
        assert "ñ" not in nombre.lower()
        assert " " not in nombre
        assert nombre.endswith(".xlsx")
        # Sigue siendo un nombre de archivo válido: solo minúsculas, dígitos y guiones.
        import re

        assert re.fullmatch(r"[a-z0-9.\-]+", nombre), nombre
